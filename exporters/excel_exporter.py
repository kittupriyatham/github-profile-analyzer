"""Excel export module for QIntern candidate screening results.

Generates a production-quality recruitment workbook with four sheets:
  1. Rankings   – summary view sorted by status then total score
  2. Detailed Scores – per-domain breakdowns with keywords and libraries
  3. Contact Info    – candidate contact details and key URLs
  4. Pipeline Log    – run metadata and aggregate statistics

Requires: openpyxl (``pip install openpyxl``)
"""

import logging
import os
from datetime import datetime
from typing import Any, Optional

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    raise ImportError(
        "openpyxl is required for Excel export. Install it with: pip install openpyxl"
    )

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Style constants
# ---------------------------------------------------------------------------

_THIN_SIDE = Side(style="thin", color="AAAAAA")
_CELL_BORDER = Border(
    left=_THIN_SIDE,
    right=_THIN_SIDE,
    top=_THIN_SIDE,
    bottom=_THIN_SIDE,
)

_HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
_HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
_HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

_STATUS_FILLS: dict[str, PatternFill] = {
    "SHORTLIST": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
    "REVIEW": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
    "REJECT": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
}

_STATUS_SORT_ORDER: dict[str, int] = {"SHORTLIST": 0, "REVIEW": 1, "REJECT": 2}

_BODY_FONT = Font(name="Calibri", size=10)
_BODY_ALIGNMENT = Alignment(vertical="center", wrap_text=True)
_NUMBER_FORMAT = "0.0"

# Domains used for column generation – order matters for readability.
_SCORE_DOMAINS = [
    "Quantum Technology",
    "ML/DL/DS/AI",
    "Software Engineering",
    "Cloud & DevOps",
    "Mathematics & Foundations",
]

_DOMAIN_SHORT_LABELS: dict[str, str] = {
    "Quantum Technology": "Quantum Score",
    "ML/DL/DS/AI": "ML/AI Score",
    "Software Engineering": "Software Score",
    "Cloud & DevOps": "Cloud Score",
    "Mathematics & Foundations": "Math Score",
}

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _safe_get(obj: Any, attr: str, default: Any = "") -> Any:
    """Return *attr* from *obj* via attribute or dict key lookup.

    Supports both dataclass-style objects and plain dicts, making the
    exporter resilient to data model changes.
    """
    if isinstance(obj, dict):
        return obj.get(attr, default)
    return getattr(obj, attr, default)


def _domain_score(candidate: Any, domain: str) -> float:
    """Extract a single domain score from a candidate record."""
    scores = _safe_get(candidate, "domain_scores", {})
    if isinstance(scores, dict):
        return float(scores.get(domain, 0))
    return 0.0


def _total_score(candidate: Any) -> float:
    """Sum all domain scores for a candidate."""
    scores = _safe_get(candidate, "domain_scores", {})
    if isinstance(scores, dict):
        return float(sum(scores.values()))
    return float(_safe_get(candidate, "total_score", 0))


def _status(candidate: Any) -> str:
    """Normalise candidate status to one of SHORTLIST / REVIEW / REJECT."""
    raw = str(_safe_get(candidate, "status", "REVIEW")).upper().strip()
    if raw in _STATUS_SORT_ORDER:
        return raw
    return "REVIEW"


def _top_items(mapping: Any, domain: str, n: int = 5) -> str:
    """Return a comma-separated string of the top *n* items for *domain*."""
    if not mapping or not isinstance(mapping, dict):
        return ""
    items = mapping.get(domain, [])
    if isinstance(items, dict):
        # Counter-style: {kw: count, …}
        sorted_items = sorted(items.items(), key=lambda kv: kv[1], reverse=True)
        return ", ".join(k for k, _ in sorted_items[:n])
    if isinstance(items, (list, tuple)):
        return ", ".join(str(i) for i in items[:n])
    return str(items)


def _join_list(value: Any) -> str:
    """Flatten a list/set/tuple into a comma-separated string."""
    if isinstance(value, (list, tuple, set, frozenset)):
        return ", ".join(str(v) for v in value)
    if value is None:
        return ""
    return str(value)


def _sort_candidates(candidates: list) -> list:
    """Sort candidates: SHORTLIST first, then REVIEW, then REJECT.

    Within each status group, sort by total score descending.
    """
    return sorted(
        candidates,
        key=lambda c: (_STATUS_SORT_ORDER.get(_status(c), 1), -_total_score(c)),
    )


# ---------------------------------------------------------------------------
# Sheet styling utilities
# ---------------------------------------------------------------------------


def _apply_header_style(ws: Any, col_count: int) -> None:
    """Style the first row as the header and freeze it."""
    for col_idx in range(1, col_count + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGNMENT
        cell.border = _CELL_BORDER
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _apply_body_borders(ws: Any, max_row: int, max_col: int) -> None:
    """Add thin borders and base font to every body cell."""
    for row_idx in range(2, max_row + 1):
        for col_idx in range(1, max_col + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = _CELL_BORDER
            cell.font = _BODY_FONT
            cell.alignment = _BODY_ALIGNMENT


def _auto_fit_columns(ws: Any, min_width: int = 10, max_width: int = 50) -> None:
    """Approximate auto-fit by scanning content length per column."""
    for col_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            try:
                val = str(cell.value) if cell.value is not None else ""
                # Use the longest single line for width calculation
                for line in val.split("\n"):
                    max_len = max(max_len, len(line))
            except Exception:
                pass
        adjusted = max(min(max_len + 3, max_width), min_width)
        ws.column_dimensions[col_letter].width = adjusted


# ---------------------------------------------------------------------------
# Sheet builders
# ---------------------------------------------------------------------------


def _build_rankings_sheet(ws: Any, candidates: list) -> None:
    """Sheet 1 – Rankings summary with color-coded status rows."""
    headers = [
        "Rank",
        "ID",
        "Name",
        "Email",
        "Status",
        "Assigned Bucket",
        "Total Score",
        "Quantum Score",
        "ML/AI Score",
        "Software Score",
        "Cloud Score",
        "Math Score",
        "Primary Domain",
        "Repos Scanned",
        "Classical GitHub",
        "Quantum GitHub",
        "Classical Deployed",
        "Quantum Deployed",
        "Reason",
    ]
    ws.append(headers)

    sorted_candidates = _sort_candidates(candidates)

    for rank, c in enumerate(sorted_candidates, start=1):
        status = _status(c)
        row = [
            rank,
            _safe_get(c, "candidate_id"),
            _safe_get(c, "full_name"),
            _safe_get(c, "email"),
            status,
            _safe_get(c, "assigned_bucket"),
            _total_score(c),
            _domain_score(c, "Quantum Technology"),
            _domain_score(c, "ML/DL/DS/AI"),
            _domain_score(c, "Software Engineering"),
            _domain_score(c, "Cloud & DevOps"),
            _domain_score(c, "Mathematics & Foundations"),
            _safe_get(c, "primary_domain"),
            _safe_get(c, "repos_scanned", 0),
            _safe_get(c, "classical_github"),
            _safe_get(c, "quantum_github"),
            _safe_get(c, "classical_deployed"),
            _safe_get(c, "quantum_deployed"),
            _safe_get(c, "status_reason"),
        ]
        ws.append(row)

        # Apply status colour fill to the entire row
        row_idx = ws.max_row
        fill = _STATUS_FILLS.get(status)
        if fill:
            for col_idx in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = fill

        # Number format for score columns (cols 7-12)
        for col_idx in range(7, 13):
            ws.cell(row=row_idx, column=col_idx).number_format = _NUMBER_FORMAT

    _apply_header_style(ws, len(headers))
    _apply_body_borders(ws, ws.max_row, len(headers))
    _auto_fit_columns(ws)


def _build_detailed_scores_sheet(ws: Any, candidates: list) -> None:
    """Sheet 2 – Detailed per-domain scores, keywords, and libraries."""
    headers = ["Name"]
    for domain in _SCORE_DOMAINS:
        label = _DOMAIN_SHORT_LABELS.get(domain, domain)
        headers.append(label)
    for domain in _SCORE_DOMAINS:
        short = _DOMAIN_SHORT_LABELS.get(domain, domain).replace(" Score", "")
        headers.append(f"Top Keywords ({short})")
    for domain in _SCORE_DOMAINS:
        short = _DOMAIN_SHORT_LABELS.get(domain, domain).replace(" Score", "")
        headers.append(f"Top Libraries ({short})")
    headers.append("Languages Used")
    ws.append(headers)

    for c in _sort_candidates(candidates):
        row: list[Any] = [_safe_get(c, "name")]

        # Domain scores
        for domain in _SCORE_DOMAINS:
            row.append(_domain_score(c, domain))

        # Top keywords per domain
        keywords_map = _safe_get(c, "domain_keywords", _safe_get(c, "keywords", {}))
        for domain in _SCORE_DOMAINS:
            row.append(_top_items(keywords_map, domain))

        # Top libraries per domain
        libraries_map = _safe_get(c, "domain_libraries", _safe_get(c, "libraries", {}))
        for domain in _SCORE_DOMAINS:
            row.append(_top_items(libraries_map, domain))

        # Languages
        languages = _safe_get(c, "languages", [])
        row.append(_join_list(languages))

        ws.append(row)

        # Number format for score columns (cols 2-6)
        row_idx = ws.max_row
        for col_idx in range(2, 2 + len(_SCORE_DOMAINS)):
            ws.cell(row=row_idx, column=col_idx).number_format = _NUMBER_FORMAT

    _apply_header_style(ws, len(headers))
    _apply_body_borders(ws, ws.max_row, len(headers))
    _auto_fit_columns(ws)


def _build_contact_info_sheet(ws: Any, candidates: list) -> None:
    """Sheet 3 – Contact details and key URLs."""
    headers = [
        "Name",
        "Email",
        "Phone",
        "Discord",
        "GitHub URL",
        "LinkedIn URL",
        "Resume URL",
        "Best Project URL",
    ]
    ws.append(headers)

    for c in _sort_candidates(candidates):
        row = [
            _safe_get(c, "name"),
            _safe_get(c, "email"),
            _safe_get(c, "phone"),
            _safe_get(c, "discord"),
            _safe_get(c, "github_url", _safe_get(c, "github")),
            _safe_get(c, "linkedin_url", _safe_get(c, "linkedin")),
            _safe_get(c, "resume_url", _safe_get(c, "resume")),
            _safe_get(c, "best_project_url", _safe_get(c, "best_project", "")),
        ]
        ws.append(row)

    _apply_header_style(ws, len(headers))
    _apply_body_borders(ws, ws.max_row, len(headers))
    _auto_fit_columns(ws)


def _build_pipeline_log_sheet(
    ws: Any,
    candidates: list,
    run_metadata: dict[str, Any] | None,
) -> None:
    """Sheet 4 – Pipeline run metadata and aggregate statistics."""
    headers = ["Metric", "Value"]
    ws.append(headers)

    # Compute aggregate counts
    status_counts: dict[str, int] = {"SHORTLIST": 0, "REVIEW": 0, "REJECT": 0}
    for c in candidates:
        s = _status(c)
        status_counts[s] = status_counts.get(s, 0) + 1

    meta = run_metadata or {}

    rows = [
        ("Run Timestamp", meta.get("timestamp", datetime.now().strftime("%Y-%m-%d %H:%M:%S"))),
        ("Total Candidates Processed", len(candidates)),
        ("Shortlisted Count", status_counts.get("SHORTLIST", 0)),
        ("Review Count", status_counts.get("REVIEW", 0)),
        ("Rejected Count", status_counts.get("REJECT", 0)),
        ("Errors Count", meta.get("errors_count", meta.get("errors", 0))),
        ("Processing Time", meta.get("processing_time", meta.get("duration", "N/A"))),
        ("Config Used", meta.get("config_used", meta.get("config", "N/A"))),
    ]

    for metric, value in rows:
        ws.append([metric, value])

    _apply_header_style(ws, len(headers))
    _apply_body_borders(ws, ws.max_row, len(headers))
    _auto_fit_columns(ws, min_width=20, max_width=60)


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------


def export_results(
    candidates: list,
    output_path: str = "qintern_results.xlsx",
    run_metadata: dict[str, Any] | None = None,
) -> str:
    """Export analysis results to a formatted Excel workbook.

    Creates a multi-sheet workbook with Rankings, Detailed Scores,
    Contact Info, and Pipeline Log sheets. Each sheet is fully formatted
    with frozen headers, auto-filters, colour-coded status rows, thin
    borders, and approximate auto-fit column widths.

    Args:
        candidates: List of candidate record objects (dataclass or dict).
            Expected attributes/keys include ``name``, ``email``,
            ``github_url``, ``status``, ``domain_scores``, etc.
        output_path: Destination file path for the ``.xlsx`` workbook.
            Parent directories are created automatically if missing.
        run_metadata: Optional dict of pipeline run info (timestamp,
            processing_time, config_used, errors_count, …).

    Returns:
        Absolute path to the created Excel file.

    Raises:
        ValueError: If *candidates* is ``None``.
        OSError: If the file cannot be written to *output_path*.
    """
    if candidates is None:
        raise ValueError("candidates list must not be None")

    if not candidates:
        logger.warning("No candidates to export – workbook will contain headers only.")

    # Ensure output directory exists
    output_dir = os.path.dirname(os.path.abspath(output_path))
    os.makedirs(output_dir, exist_ok=True)

    wb = Workbook()

    # --- Sheet 1: Rankings ---
    ws_rankings = wb.active
    ws_rankings.title = "Rankings"
    _build_rankings_sheet(ws_rankings, candidates)

    # --- Sheet 2: Detailed Scores ---
    ws_detailed = wb.create_sheet(title="Detailed Scores")
    _build_detailed_scores_sheet(ws_detailed, candidates)

    # --- Sheet 3: Contact Info ---
    ws_contact = wb.create_sheet(title="Contact Info")
    _build_contact_info_sheet(ws_contact, candidates)

    # --- Sheet 4: Pipeline Log ---
    ws_log = wb.create_sheet(title="Pipeline Log")
    _build_pipeline_log_sheet(ws_log, candidates, run_metadata)

    # Persist workbook
    try:
        wb.save(output_path)
    except OSError as exc:
        logger.error(f"Failed to save workbook to {output_path}: {exc}")
        raise

    abs_path = os.path.abspath(output_path)
    logger.info(f"Workbook saved to {abs_path}")
    return abs_path
